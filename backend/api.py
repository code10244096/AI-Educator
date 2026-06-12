from fastapi import APIRouter, UploadFile, File, Depends, HTTPException, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, delete
from typing import List, Optional
import json
import os
from datetime import datetime
import aiofiles

from database import get_db, init_db
from models import (
    HomeworkAssignment, HomeworkSubmission, WrongQuestion,
    LessonPlan, User, ClassInfo, QuestionBank, QuestionVariant
)
from ai_client import ai_client
from rag_retriever import QuestionRetriever
from config import settings
from homework_dataset import list_dataset_homeworks, get_dataset_homework, get_dataset_file_path
from class_service import (
    seed_initial_data,
    get_class_list,
    get_homework_list,
    get_homework_by_id,
    get_student_submissions,
    get_homework_stats,
    get_grading_tasks,
    get_alert_students,
    update_assignment_status,
    sync_wrong_questions,
    resolve_class_id,
    get_assignment_record,
)

router = APIRouter()


def ensure_upload_dir():
    """确保上传目录存在"""
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)


async def _save_grading_result(
    db: AsyncSession,
    grading_result: dict,
    ocr_result: str,
    file_paths: List[str],
    *,
    assignment_id: Optional[int] = None,
    submission_id: Optional[int] = None,
    student_name: Optional[str] = None,
    dataset_file_id: Optional[int] = None,
    is_test_data: bool = False,
    subject: str = "数学",
    sync_notebook: bool = True,
) -> HomeworkSubmission:
    """保存或更新批改结果，并同步班级/错题本状态"""
    if submission_id:
        result = await db.execute(
            select(HomeworkSubmission).where(HomeworkSubmission.id == submission_id)
        )
        submission = result.scalar_one_or_none()
        if not submission:
            raise HTTPException(status_code=404, detail="提交记录不存在")
        submission.ocr_result = ocr_result
        submission.grading_result = json.dumps(grading_result, ensure_ascii=False)
        submission.wrong_count = grading_result.get("wrong_count", 0)
        submission.score = grading_result.get("score", 0)
        submission.status = "completed"
        submission.grading_status = "已批改"
        submission.image_paths = json.dumps(file_paths)
        if dataset_file_id:
            submission.dataset_file_id = dataset_file_id
    else:
        submission = HomeworkSubmission(
            assignment_id=assignment_id,
            student_name=student_name,
            dataset_file_id=dataset_file_id,
            is_test_data=is_test_data,
            image_paths=json.dumps(file_paths),
            ocr_result=ocr_result,
            grading_result=json.dumps(grading_result, ensure_ascii=False),
            wrong_count=grading_result.get("wrong_count", 0),
            score=grading_result.get("score", 0),
            status="completed",
            grading_status="已批改",
            submit_time=datetime.now().strftime("%Y-%m-%d %H:%M"),
            file_count=len(file_paths),
        )
        db.add(submission)

    await db.flush()

    if sync_notebook and student_name:
        await sync_wrong_questions(db, grading_result, student_name, subject)

    if submission.assignment_id:
        await update_assignment_status(db, submission.assignment_id)

    await db.commit()
    await db.refresh(submission)
    return submission


@router.post("/grader/upload")
async def upload_homework(
    files: List[UploadFile] = File(...),
    reference_answer: Optional[str] = Form(None),
    subject: str = Form("数学"),
    assignment_id: Optional[int] = Form(None),
    submission_id: Optional[int] = Form(None),
    student_name: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db)
):
    """上传作业文件并批改 - 支持图片、PDF、Word、Excel、TXT、MD等多种格式"""
    ensure_upload_dir()
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_paths = []
    text_contents = []
    
    for idx, file in enumerate(files):
        filename = f"{timestamp}_{idx}_{file.filename}"
        filepath = os.path.join(settings.UPLOAD_DIR, filename)
        
        async with aiofiles.open(filepath, "wb") as out_file:
            content = await file.read()
            await out_file.write(content)
        
        file_paths.append(filepath)
        file_ext = file.filename.split('.')[-1].lower() if '.' in file.filename else ''
        
        try:
            if file_ext in ['png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp']:
                ocr_text = await ai_client.ocr_image(filepath)
                text_contents.append(ocr_text)
            elif file_ext == 'pdf':
                try:
                    import fitz
                    doc = fitz.open(filepath)
                    text_parts = []
                    for page in doc:
                        text_parts.append(page.get_text())
                    text_contents.append("\n".join(text_parts))
                    doc.close()
                except ImportError:
                    text_contents.append(await ai_client.ocr_image(filepath))
                except Exception:
                    text_contents.append(await ai_client.ocr_image(filepath))
            elif file_ext in ['doc', 'docx']:
                try:
                    from docx import Document
                    doc = Document(filepath)
                    text_parts = [para.text for para in doc.paragraphs]
                    text_contents.append("\n".join(text_parts))
                except ImportError:
                    text_contents.append(await ai_client.ocr_image(filepath))
                except Exception:
                    text_contents.append(await ai_client.ocr_image(filepath))
            elif file_ext in ['xls', 'xlsx']:
                try:
                    import pandas as pd
                    df = pd.read_excel(filepath)
                    text_contents.append(df.to_string())
                except ImportError:
                    text_contents.append(await ai_client.ocr_image(filepath))
                except Exception:
                    text_contents.append(await ai_client.ocr_image(filepath))
            elif file_ext in ['txt', 'md']:
                async with aiofiles.open(filepath, "r", encoding="utf-8") as f:
                    text_contents.append(await f.read())
            else:
                text_contents.append(await ai_client.ocr_image(filepath))
        except Exception as e:
            text_contents.append(f"[文件解析失败: {file.filename}]")
    
    full_text_result = "\n\n".join(text_contents)
    
    if assignment_id and not reference_answer:
        assign_result = await db.execute(
            select(HomeworkAssignment).where(HomeworkAssignment.id == assignment_id)
        )
        assignment = assign_result.scalar_one_or_none()
        if assignment and assignment.reference_answer:
            reference_answer = assignment.reference_answer

    grading_result = await ai_client.grade_homework(
        ocr_result=full_text_result,
        reference_answer=reference_answer,
        subject=subject
    )

    submission = await _save_grading_result(
        db,
        grading_result,
        full_text_result,
        file_paths,
        assignment_id=assignment_id,
        submission_id=submission_id,
        student_name=student_name,
        subject=subject,
    )

    return {
        "submission_id": submission.id,
        "assignment_id": submission.assignment_id,
        "student_name": submission.student_name,
        "ocr_result": full_text_result,
        "grading_result": grading_result,
        "image_count": len(file_paths)
    }


@router.get("/grader/{submission_id}")
async def get_grading_result(submission_id: int, db: AsyncSession = Depends(get_db)):
    """获取批改结果"""
    result = await db.execute(
        select(HomeworkSubmission).where(HomeworkSubmission.id == submission_id)
    )
    submission = result.scalar_one_or_none()
    
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    
    return {
        "id": submission.id,
        "assignment_id": submission.assignment_id,
        "student_name": submission.student_name,
        "ocr_result": submission.ocr_result,
        "grading_result": json.loads(submission.grading_result) if submission.grading_result else {},
        "wrong_count": submission.wrong_count,
        "score": submission.score,
        "status": submission.status,
        "grading_status": submission.grading_status,
        "created_at": submission.created_at
    }


@router.get("/homework/dataset")
async def get_homework_dataset_list():
    """获取 dataset 测试集作业列表"""
    return {"items": list_dataset_homeworks(), "total": len(list_dataset_homeworks())}


@router.get("/homework/dataset/{file_id}")
async def get_homework_dataset_detail(file_id: int):
    """获取 dataset 测试集作业详情（含学生作答与参考答案）"""
    data = get_dataset_homework(file_id=file_id)
    if not data:
        raise HTTPException(status_code=404, detail="测试作业不存在")
    return data


@router.post("/grader/upload-dataset/{file_id}")
async def upload_dataset_homework(
    file_id: int,
    subject: str = Form("数学"),
    assignment_id: Optional[int] = Form(None),
    submission_id: Optional[int] = Form(None),
    student_name: Optional[str] = Form(None),
    class_slug: Optional[str] = Form(None),
    homework_id: Optional[int] = Form(None),
    db: AsyncSession = Depends(get_db)
):
    """使用 dataset 测试集文件直接批改（调试用）"""
    data = get_dataset_homework(file_id=file_id)
    if not data:
        raise HTTPException(status_code=404, detail="测试作业不存在")

    filepath = get_dataset_file_path(data["filename"])
    if not filepath:
        raise HTTPException(status_code=404, detail="测试文件不存在")

    full_text_result = data["full_content"]
    reference_answer = data["reference_answer"]

    if not assignment_id and class_slug and homework_id:
        from class_service import get_assignment_record
        assignment = await get_assignment_record(db, class_slug, homework_id)
        if assignment:
            assignment_id = assignment.id
            if not reference_answer:
                reference_answer = assignment.reference_answer

    if submission_id and not student_name:
        sub_result = await db.execute(
            select(HomeworkSubmission).where(HomeworkSubmission.id == submission_id)
        )
        existing = sub_result.scalar_one_or_none()
        if existing:
            student_name = existing.student_name

    grading_result = await ai_client.grade_homework(
        ocr_result=full_text_result,
        reference_answer=reference_answer,
        subject=subject
    )

    submission = await _save_grading_result(
        db,
        grading_result,
        full_text_result,
        [filepath],
        assignment_id=assignment_id,
        submission_id=submission_id,
        student_name=student_name or data["title"],
        dataset_file_id=file_id,
        is_test_data=True,
        subject=subject,
    )

    return {
        "submission_id": submission.id,
        "assignment_id": submission.assignment_id,
        "student_name": submission.student_name,
        "dataset_title": data["title"],
        "dataset_filename": data["filename"],
        "ocr_result": full_text_result,
        "grading_result": grading_result,
        "image_count": 1
    }


@router.post("/notebook/upload")
async def upload_wrong_question(
    file: UploadFile = File(...),
    knowledge_point: str = Form(...),
    subject: str = Form("数学"),
    db: AsyncSession = Depends(get_db)
):
    """录入错题 - 支持图片、PDF、Word、Excel、TXT、MD等多种格式"""
    ensure_upload_dir()
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{timestamp}_{file.filename}"
    filepath = os.path.join(settings.UPLOAD_DIR, filename)
    
    async with aiofiles.open(filepath, "wb") as out_file:
        content = await file.read()
        await out_file.write(content)
    
    file_ext = file.filename.split('.')[-1].lower() if '.' in file.filename else ''
    text_content = ""
    
    try:
        if file_ext in ['png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp']:
            ocr_result = await ai_client.ocr_image(filepath)
            text_content = ocr_result
        elif file_ext == 'pdf':
            try:
                import fitz
                doc = fitz.open(filepath)
                text_parts = []
                for page in doc:
                    text_parts.append(page.get_text())
                text_content = "\n".join(text_parts)
                doc.close()
            except ImportError:
                text_content = await ai_client.ocr_image(filepath)
            except Exception:
                text_content = await ai_client.ocr_image(filepath)
        elif file_ext in ['doc', 'docx']:
            try:
                from docx import Document
                doc = Document(filepath)
                text_parts = [para.text for para in doc.paragraphs]
                text_content = "\n".join(text_parts)
            except ImportError:
                text_content = await ai_client.ocr_image(filepath)
            except Exception:
                text_content = await ai_client.ocr_image(filepath)
        elif file_ext in ['xls', 'xlsx']:
            try:
                import pandas as pd
                df = pd.read_excel(filepath)
                text_content = df.to_string()
            except ImportError:
                text_content = await ai_client.ocr_image(filepath)
            except Exception:
                text_content = await ai_client.ocr_image(filepath)
        elif file_ext in ['txt', 'md']:
            async with aiofiles.open(filepath, "r", encoding="utf-8") as f:
                text_content = await f.read()
        else:
            return {
                "error": "这个文件格式我不太认识呢，试试图片、PDF、Word、Excel 或 TXT 文件吧~",
                "error_type": "invalid_format",
                "questions": []
            }
    except Exception as e:
        return {
            "error": f"解析遇到了一点小麻烦：{str(e)}",
            "error_type": "parse_failed",
            "questions": []
        }
    
    if not text_content or len(text_content.strip()) < 10:
        return {
            "error": "哎呀，这个文件好像有点'害羞'，什么都没解析出来呢~ 换个文件试试看？",
            "error_type": "empty",
            "questions": []
        }
    
    question_text = text_content.split("答案")[0] if "答案" in text_content else text_content
    user_answer = ""
    correct_answer = ""
    
    if "答案" in text_content:
        parts = text_content.split("答案")
        if len(parts) > 1:
            correct_answer = parts[1]
    
    try:
        variant_questions = await ai_client.generate_variant_questions(
            question_text=question_text,
            knowledge_point=knowledge_point,
            count=3
        )
    except:
        variant_questions = []
    
    wrong_question = WrongQuestion(
        user_id=1,
        question_text=question_text,
        user_answer=user_answer,
        correct_answer=correct_answer,
        knowledge_point=knowledge_point,
        subject=subject,
        variant_questions=json.dumps(variant_questions, ensure_ascii=False),
        image_path=filepath
    )
    
    db.add(wrong_question)
    await db.commit()
    await db.refresh(wrong_question)
    
    return {
        "id": wrong_question.id,
        "questions": [{
            "id": wrong_question.id,
            "question_text": question_text,
            "knowledge_point": knowledge_point,
            "variant_questions": variant_questions,
            "image_path": filepath
        }],
        "knowledge_point": knowledge_point,
        "variant_questions": variant_questions,
        "image_path": filepath
    }


@router.get("/notebook/list")
async def get_wrong_questions(
    knowledge_point: Optional[str] = None,
    subject: Optional[str] = None,
    limit: int = 20,
    offset: int = 0,
    db: AsyncSession = Depends(get_db)
):
    """获取错题列表"""
    query = select(WrongQuestion)
    
    if knowledge_point:
        query = query.where(WrongQuestion.knowledge_point.like(f"%{knowledge_point}%"))
    if subject:
        query = query.where(WrongQuestion.subject == subject)
    
    query = query.order_by(WrongQuestion.error_date.desc())
    query = query.offset(offset).limit(limit)
    
    result = await db.execute(query)
    questions = result.scalars().all()
    
    return [
        {
            "id": q.id,
            "question_text": q.question_text,
            "user_answer": q.user_answer,
            "correct_answer": q.correct_answer,
            "knowledge_point": q.knowledge_point,
            "subject": q.subject,
            "error_date": q.error_date,
            "is_mastered": q.is_mastered,
            "variant_questions": json.loads(q.variant_questions) if q.variant_questions else []
        }
        for q in questions
    ]


@router.post("/notebook/{question_id}/mastered")
async def mark_as_mastered(question_id: int, db: AsyncSession = Depends(get_db)):
    """标记为已掌握"""
    result = await db.execute(
        select(WrongQuestion).where(WrongQuestion.id == question_id)
    )
    question = result.scalar_one_or_none()
    
    if not question:
        raise HTTPException(status_code=404, detail="错题不存在")
    
    question.is_mastered = True
    await db.commit()
    
    return {"message": "已标记为已掌握"}


@router.post("/lessonplan/generate")
async def generate_lesson_plan(
    title: str = Form(...),
    period: str = Form("1 课时"),
    student_level: str = Form("中等"),
    requirements: str = Form(""),
    db: AsyncSession = Depends(get_db)
):
    """生成教案 - 集成题库RAG检索"""
    # 从题库检索相关题目
    try:
        retriever = QuestionRetriever()
        related_questions = await retriever.keyword_search(
            db=db,
            query_text=title,
            subject="数学",
            education_level="高中",
            limit=8
        )
        
        # 格式化为 RAG 上下文
        question_bank_context = retriever.format_for_rag(related_questions)
    except Exception as e:
        print(f"题库检索失败: {e}")
        question_bank_context = ""
    
    try:
        content = await ai_client.generate_lesson_plan(
            topic=title,
            period=period,
            student_level=student_level,
            requirements=requirements,
            question_bank_context=question_bank_context
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"教案生成失败：{str(e)}")
    
    # 保存教案
    lesson_plan = LessonPlan(
        teacher_id=1,  # TODO: 从登录用户获取
        title=title,
        period=period,
        student_level=student_level,
        requirements=requirements,
        content=content
    )
    
    db.add(lesson_plan)
    await db.commit()
    await db.refresh(lesson_plan)
    
    return {
        "id": lesson_plan.id,
        "title": lesson_plan.title,
        "content": content,
        "created_at": lesson_plan.created_at,
        "retrieved_questions_count": len(related_questions) if related_questions else 0
    }


@router.get("/lessonplan/{plan_id}")
async def get_lesson_plan(plan_id: int, db: AsyncSession = Depends(get_db)):
    """获取教案详情"""
    result = await db.execute(
        select(LessonPlan).where(LessonPlan.id == plan_id)
    )
    plan = result.scalar_one_or_none()
    
    if not plan:
        raise HTTPException(status_code=404, detail="教案不存在")
    
    return {
        "id": plan.id,
        "title": plan.title,
        "period": plan.period,
        "student_level": plan.student_level,
        "requirements": plan.requirements,
        "content": plan.content,
        "created_at": plan.created_at,
        "updated_at": plan.updated_at
    }


@router.get("/class/list")
async def list_classes(db: AsyncSession = Depends(get_db)):
    """获取班级列表"""
    return {"items": await get_class_list(db)}


@router.get("/class/{class_slug}/homework")
async def list_class_homework(class_slug: str, db: AsyncSession = Depends(get_db)):
    """获取班级作业列表"""
    try:
        return {"items": await get_homework_list(db, class_slug)}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/class/{class_slug}/homework/{homework_id}")
async def get_class_homework_detail(
    class_slug: str, homework_id: int, db: AsyncSession = Depends(get_db)
):
    """获取班级作业详情"""
    try:
        homework = await get_homework_by_id(db, class_slug, homework_id)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    if not homework:
        raise HTTPException(status_code=404, detail="作业不存在")
    return homework


@router.get("/class/{class_slug}/homework/{homework_id}/submissions")
async def list_homework_submissions(
    class_slug: str, homework_id: int, db: AsyncSession = Depends(get_db)
):
    """获取作业学生提交列表"""
    try:
        return {"items": await get_student_submissions(db, class_slug, homework_id)}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/class/{class_slug}/homework-stats")
async def class_homework_stats(class_slug: str, db: AsyncSession = Depends(get_db)):
    """获取班级作业看板统计"""
    try:
        return await get_homework_stats(db, class_slug)
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/class/{class_slug}/grading-tasks")
async def class_grading_tasks(class_slug: str, db: AsyncSession = Depends(get_db)):
    """获取班级待批改任务"""
    try:
        return {"items": await get_grading_tasks(db, class_slug)}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/class/{class_slug}/alert-students")
async def class_alert_students(class_slug: str, db: AsyncSession = Depends(get_db)):
    """获取预警学生"""
    try:
        return {"items": await get_alert_students(db, class_slug)}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.delete("/class/{class_slug}/homework/{homework_id}/submissions/pending")
async def delete_pending_submissions(
    class_slug: str, homework_id: int, db: AsyncSession = Depends(get_db)
):
    """批量删除待批改的作业提交记录"""
    try:
        assignment = await get_assignment_record(db, class_slug, homework_id)
        if not assignment:
            raise HTTPException(status_code=404, detail="作业不存在")
        
        # 删除所有待批改状态的提交记录
        result = await db.execute(
            delete(HomeworkSubmission)
            .where(HomeworkSubmission.assignment_id == assignment.id)
            .where(HomeworkSubmission.grading_status == "待批改")
        )
        
        await db.commit()
        
        return {
            "message": f"成功删除 {result.rowcount} 条待批改作业记录",
            "deleted_count": result.rowcount
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.delete("/class/{class_slug}/homework/{homework_id}/submissions/keep-first/{keep_count}")
async def keep_first_n_submissions(
    class_slug: str, homework_id: int, keep_count: int, db: AsyncSession = Depends(get_db)
):
    """保留前N条学生提交记录，删除其余所有记录"""
    try:
        assignment = await get_assignment_record(db, class_slug, homework_id)
        if not assignment:
            raise HTTPException(status_code=404, detail="作业不存在")
        
        # 获取前keep_count条记录的ID
        select_result = await db.execute(
            select(HomeworkSubmission.id)
            .where(HomeworkSubmission.assignment_id == assignment.id)
            .order_by(HomeworkSubmission.id)
            .limit(keep_count)
        )
        keep_ids = [row[0] for row in select_result.all()]
        
        # 删除不在保留列表中的记录
        result = await db.execute(
            delete(HomeworkSubmission)
            .where(HomeworkSubmission.assignment_id == assignment.id)
            .where(HomeworkSubmission.id.not_in(keep_ids if keep_ids else [-1]))
        )
        
        await db.commit()
        
        return {
            "message": f"成功保留前 {keep_count} 条记录，删除 {result.rowcount} 条记录",
            "deleted_count": result.rowcount,
            "kept_count": keep_count
        }
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))


@router.get("/class/stats")
async def get_class_stats(
    class_slug: Optional[str] = None,
    db: AsyncSession = Depends(get_db)
):
    """获取班级统计信息（支持按班级筛选）"""
    query = select(HomeworkSubmission)
    if class_slug:
        try:
            class_id = resolve_class_id(class_slug)
        except ValueError as e:
            raise HTTPException(status_code=404, detail=str(e))
        assign_result = await db.execute(
            select(HomeworkAssignment.id).where(HomeworkAssignment.class_id == class_id)
        )
        assignment_ids = [row[0] for row in assign_result.all()]
        if assignment_ids:
            query = query.where(HomeworkSubmission.assignment_id.in_(assignment_ids))
        else:
            return {
                "total_students": 0,
                "average_wrong_count": 0,
                "common_wrong_questions": []
            }

    result = await db.execute(query)
    submissions = result.scalars().all()

    if not submissions:
        return {
            "total_students": 0,
            "average_wrong_count": 0,
            "common_wrong_questions": []
        }

    total_students = len(set(s.student_name for s in submissions if s.student_name))
    total_wrong = sum(s.wrong_count or 0 for s in submissions)
    avg_wrong = total_wrong / len(submissions) if submissions else 0

    wrong_questions = {}
    for sub in submissions:
        grading = json.loads(sub.grading_result) if sub.grading_result else {}
        for q in grading.get("questions", []):
            if not q.get("is_correct", True):
                q_text = q.get("question_text", "")[:50]
                wrong_questions[q_text] = wrong_questions.get(q_text, 0) + 1

    common_wrong = sorted(wrong_questions.items(), key=lambda x: x[1], reverse=True)[:5]

    return {
        "total_students": total_students,
        "average_wrong_count": round(avg_wrong, 1),
        "common_wrong_questions": [
            {"question": q, "count": c}
            for q, c in common_wrong
        ]
    }


@router.get("/tasks/all")
async def get_all_grading_tasks(db: AsyncSession = Depends(get_db)):
    """聚合所有班级的批改任务（供我的任务页使用）"""
    all_tasks = []
    for slug in ["class1", "class2", "class3"]:
        try:
            tasks = await get_grading_tasks(db, slug)
            all_tasks.extend(tasks)
        except ValueError:
            continue
    return {"items": all_tasks}


@router.on_event("startup")
async def startup_event():
    """启动时初始化数据库并播种班级作业数据"""
    from database import AsyncSessionLocal
    await init_db()
    async with AsyncSessionLocal() as db:
        await seed_initial_data(db)


# ==================== 题库管理 API ====================

@router.post("/questionbank/add")
async def add_question(
    question_text: str = Form(...),
    answer: str = Form(...),
    solution: str = Form(""),
    question_type: str = Form(...),
    subject: str = Form(...),
    education_level: str = Form(...),
    exam_type: str = Form(""),
    year: int = Form(None),
    region: str = Form(""),
    knowledge_points: str = Form(""),  # JSON 字符串
    difficulty: int = Form(3),
    score: float = Form(None),
    source_url: str = Form(""),
    teaching_tips: str = Form(""),
    common_mistakes: str = Form(""),
    db: AsyncSession = Depends(get_db)
):
    """添加题目到题库"""
    question = QuestionBank(
        question_text=question_text,
        answer=answer,
        solution=solution,
        question_type=question_type,
        subject=subject,
        education_level=education_level,
        exam_type=exam_type,
        year=year,
        region=region,
        knowledge_points=knowledge_points,
        difficulty=difficulty,
        score=score,
        source_url=source_url,
        teaching_tips=teaching_tips,
        common_mistakes=common_mistakes,
        is_verified=True
    )
    
    db.add(question)
    await db.commit()
    await db.refresh(question)
    
    return {
        "id": question.id,
        "message": "题目添加成功"
    }


@router.post("/questionbank/batch-add")
async def batch_add_questions(
    questions: List[dict],
    db: AsyncSession = Depends(get_db)
):
    """批量添加题目（用于爬虫导入）"""
    added_count = 0
    
    for q_data in questions:
        question = QuestionBank(
            question_text=q_data.get("question_text"),
            answer=q_data.get("answer"),
            solution=q_data.get("solution", ""),
            question_type=q_data.get("question_type"),
            subject=q_data.get("subject"),
            education_level=q_data.get("education_level"),
            exam_type=q_data.get("exam_type", ""),
            year=q_data.get("year"),
            region=q_data.get("region", ""),
            knowledge_points=json.dumps(q_data.get("knowledge_points", []), ensure_ascii=False),
            difficulty=q_data.get("difficulty", 3),
            score=q_data.get("score"),
            source_url=q_data.get("source_url", ""),
            teaching_tips=q_data.get("teaching_tips", ""),
            common_mistakes=q_data.get("common_mistakes", ""),
            is_verified=False  # 爬虫导入的需要审核
        )
        
        db.add(question)
        added_count += 1
    
    await db.commit()
    
    return {
        "added_count": added_count,
        "message": f"成功添加 {added_count} 道题目"
    }


@router.get("/questionbank/list")
async def get_question_list(
    subject: Optional[str] = None,
    education_level: Optional[str] = None,
    exam_type: Optional[str] = None,
    year: Optional[int] = None,
    question_type: Optional[str] = None,
    knowledge_point: Optional[str] = None,
    difficulty: Optional[int] = None,
    is_verified: Optional[bool] = None,
    limit: int = 20,
    offset: int = 0,
    db: AsyncSession = Depends(get_db)
):
    """获取题目列表（支持多条件筛选）"""
    query = select(QuestionBank)
    
    if subject:
        query = query.where(QuestionBank.subject == subject)
    if education_level:
        query = query.where(QuestionBank.education_level == education_level)
    if exam_type:
        query = query.where(QuestionBank.exam_type == exam_type)
    if year:
        query = query.where(QuestionBank.year == year)
    if question_type:
        query = query.where(QuestionBank.question_type == question_type)
    if knowledge_point:
        query = query.where(QuestionBank.knowledge_points.like(f"%{knowledge_point}%"))
    if difficulty:
        query = query.where(QuestionBank.difficulty == difficulty)
    if is_verified is not None:
        query = query.where(QuestionBank.is_verified == is_verified)
    
    query = query.order_by(QuestionBank.created_at.desc())
    query = query.offset(offset).limit(limit)
    
    result = await db.execute(query)
    questions = result.scalars().all()
    
    return [
        {
            "id": q.id,
            "question_text": q.question_text,
            "answer": q.answer,
            "solution": q.solution,
            "question_type": q.question_type,
            "subject": q.subject,
            "education_level": q.education_level,
            "exam_type": q.exam_type,
            "year": q.year,
            "region": q.region,
            "knowledge_points": json.loads(q.knowledge_points) if q.knowledge_points else [],
            "difficulty": q.difficulty,
            "score": q.score,
            "teaching_tips": q.teaching_tips,
            "common_mistakes": q.common_mistakes,
            "is_verified": q.is_verified,
            "created_at": q.created_at
        }
        for q in questions
    ]


@router.get("/questionbank/{question_id}")
async def get_question_detail(question_id: int, db: AsyncSession = Depends(get_db)):
    """获取题目详情（包含变式题）"""
    result = await db.execute(
        select(QuestionBank).where(QuestionBank.id == question_id)
    )
    question = result.scalar_one_or_none()
    
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")
    
    # 获取变式题
    variants_result = await db.execute(
        select(QuestionVariant).where(QuestionVariant.original_id == question_id)
    )
    variants = variants_result.scalars().all()
    
    return {
        "id": question.id,
        "question_text": question.question_text,
        "answer": question.answer,
        "solution": question.solution,
        "question_type": question.question_type,
        "subject": question.subject,
        "education_level": question.education_level,
        "exam_type": question.exam_type,
        "year": question.year,
        "region": question.region,
        "knowledge_points": json.loads(question.knowledge_points) if question.knowledge_points else [],
        "difficulty": question.difficulty,
        "score": question.score,
        "source_url": question.source_url,
        "teaching_tips": question.teaching_tips,
        "common_mistakes": question.common_mistakes,
        "is_verified": question.is_verified,
        "variants": [
            {
                "id": v.id,
                "question_text": v.question_text,
                "answer": v.answer,
                "solution": v.solution,
                "variant_type": v.variant_type
            }
            for v in variants
        ]
    }


@router.post("/questionbank/search")
async def search_questions(
    keyword: str = Form(...),
    subject: Optional[str] = Form(None),
    education_level: Optional[str] = Form(None),
    limit: int = Form(10),
    db: AsyncSession = Depends(get_db)
):
    """搜索题目（用于 RAG 检索）"""
    # 简单关键词搜索（后续可升级为向量检索）
    query = select(QuestionBank).where(
        QuestionBank.question_text.like(f"%{keyword}%")
    )
    
    if subject:
        query = query.where(QuestionBank.subject == subject)
    if education_level:
        query = query.where(QuestionBank.education_level == education_level)
    
    query = query.limit(limit)
    
    result = await db.execute(query)
    questions = result.scalars().all()
    
    return [
        {
            "id": q.id,
            "question_text": q.question_text,
            "answer": q.answer,
            "solution": q.solution,
            "subject": q.subject,
            "education_level": q.education_level,
            "difficulty": q.difficulty
        }
        for q in questions
    ]


@router.get("/questionbank/stats")
async def get_question_stats(db: AsyncSession = Depends(get_db)):
    """获取题库统计信息"""
    # 按学科统计
    subject_result = await db.execute(
        select(QuestionBank.subject, func.count(QuestionBank.id))
        .group_by(QuestionBank.subject)
    )
    subject_stats = dict(subject_result.all())
    
    # 按学龄统计
    level_result = await db.execute(
        select(QuestionBank.education_level, func.count(QuestionBank.id))
        .group_by(QuestionBank.education_level)
    )
    level_stats = dict(level_result.all())
    
    # 总数
    total_result = await db.execute(select(func.count(QuestionBank.id)))
    total = total_result.scalar()
    
    return {
        "total": total,
        "by_subject": subject_stats,
        "by_education_level": level_stats
    }
