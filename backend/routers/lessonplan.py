"""教案：后台生成、列表、查看、编辑、删除、导出"""
import io
import re
from typing import Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, Form, HTTPException, Query, Request
from fastapi.exceptions import RequestValidationError
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from auth import current_user
from database import get_db
from jobs import run_lesson_plan_job, spawn
from models import LessonPlan, User

router = APIRouter()


class LessonPlanUpdate(BaseModel):
    title: Optional[str] = None
    content: Optional[str] = None
    period: Optional[str] = None
    student_level: Optional[str] = None
    requirements: Optional[str] = None


def _plan_status(plan: LessonPlan) -> str:
    return plan.status or "completed"


def _plan_to_dict(plan: LessonPlan) -> dict:
    return {
        "id": plan.id,
        "title": plan.title,
        "period": plan.period,
        "student_level": plan.student_level,
        "requirements": plan.requirements,
        "content": plan.content or "",
        "status": _plan_status(plan),
        "progress_stage": plan.progress_stage,
        "error_message": plan.error_message,
        "retrieved_questions_count": plan.retrieved_questions_count or 0,
        "created_at": plan.created_at,
        "updated_at": plan.updated_at,
    }


def _plan_summary(plan: LessonPlan) -> dict:
    content = plan.content or ""
    # 去掉 Markdown 标记，但保留正文里的普通连字符（如 QA-教案-一次函数）
    preview = re.sub(r"[#>*`\|$]+", " ", content)
    preview = re.sub(r"(?<!\S)-+(?!\S)", " ", preview)
    preview = re.sub(r"\s+", " ", preview).strip()[:120]
    data = _plan_to_dict(plan)
    data.pop("content")
    data["preview"] = preview
    data["content_length"] = len(content)
    return data


async def _get_plan(db: AsyncSession, plan_id: int, teacher_id: int) -> LessonPlan:
    """取当前教师的教案；不存在或属于其他教师一律 404"""
    plan = await db.get(LessonPlan, plan_id)
    if not plan or plan.teacher_id != teacher_id:
        raise HTTPException(status_code=404, detail="教案不存在")
    return plan


async def _run_or_spawn(db: AsyncSession, plan: LessonPlan, wait: bool) -> dict:
    plan_id = plan.id
    teacher_id = plan.teacher_id
    if wait:
        await run_lesson_plan_job(plan_id)
        db.expire_all()
        plan = await _get_plan(db, plan_id, teacher_id)
        if _plan_status(plan) == "failed":
            raise HTTPException(status_code=502, detail=plan.error_message or "教案生成失败")
    else:
        spawn(run_lesson_plan_job(plan_id))
        await db.refresh(plan)
    return _plan_to_dict(plan)


@router.post("/lessonplan/generate")
async def generate_lesson_plan(
    request: Request,
    period: str = Form("1 课时"),
    student_level: str = Form("中等"),
    requirements: str = Form(""),
    wait: bool = Form(False),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """
    创建教案生成任务（集成题库 RAG 检索）。
    默认立即返回 status=processing，前端轮询 GET /lessonplan/{id}；wait=true 时同步等待生成完成。
    """
    # Form(...) 会把空字符串当成缺字段，在进入本函数前就返回 422。
    # 这里直接读表单：未传 title 仍是 422，title="" 与纯空白走下面的中文 400。
    form = await request.form()
    if "title" not in form:
        raise RequestValidationError([{
            "type": "missing",
            "loc": ("body", "title"),
            "msg": "Field required",
            "input": None,
        }])
    raw_title = form.get("title")
    title = raw_title.strip() if isinstance(raw_title, str) else ""
    if not title:
        raise HTTPException(status_code=400, detail="请输入课题名称")

    plan = LessonPlan(
        teacher_id=user.id,
        title=title[:200],
        period=period,
        student_level=student_level,
        requirements=requirements,
        content="",
        status="processing",
        progress_stage="排队中",
    )
    db.add(plan)
    await db.commit()
    await db.refresh(plan)
    return await _run_or_spawn(db, plan, wait)


@router.get("/lessonplan/list")
async def list_lesson_plans(
    keyword: Optional[str] = None,
    status: Optional[str] = None,
    limit: int = Query(20, ge=1, le=200),
    offset: int = Query(0, ge=0),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """已保存的教案列表（按创建时间倒序）"""
    query = select(LessonPlan).where(LessonPlan.teacher_id == user.id)
    if keyword:
        query = query.where(or_(
            LessonPlan.title.like(f"%{keyword}%"),
            LessonPlan.requirements.like(f"%{keyword}%"),
        ))
    if status:
        if status == "completed":
            query = query.where(or_(LessonPlan.status == "completed", LessonPlan.status.is_(None)))
        else:
            query = query.where(LessonPlan.status == status)
    total = await db.scalar(select(func.count()).select_from(query.subquery()))
    rows = (await db.execute(
        query.order_by(LessonPlan.id.desc()).offset(offset).limit(limit)
    )).scalars().all()
    return {"items": [_plan_summary(p) for p in rows], "total": total or 0}


@router.get("/lessonplan/{plan_id}")
async def get_lesson_plan(plan_id: int, db: AsyncSession = Depends(get_db),
                          user: User = Depends(current_user)):
    """获取教案详情 / 生成任务状态"""
    return _plan_to_dict(await _get_plan(db, plan_id, user.id))


@router.put("/lessonplan/{plan_id}")
async def update_lesson_plan(plan_id: int, body: LessonPlanUpdate, db: AsyncSession = Depends(get_db),
                             user: User = Depends(current_user)):
    """编辑并保存教案"""
    plan = await _get_plan(db, plan_id, user.id)
    if _plan_status(plan) == "processing":
        raise HTTPException(status_code=409, detail="教案正在生成中，完成后才能编辑")
    updates = body.model_dump(exclude_none=True)
    if "title" in updates:
        updates["title"] = updates["title"].strip()
        if not updates["title"]:
            raise HTTPException(status_code=400, detail="课题名称不能为空")
    for key, value in updates.items():
        setattr(plan, key, value)
    if "content" in updates and _plan_status(plan) == "failed" and updates["content"].strip():
        plan.status = "completed"
        plan.error_message = None
    await db.commit()
    await db.refresh(plan)
    return _plan_to_dict(plan)


@router.delete("/lessonplan/{plan_id}")
async def delete_lesson_plan(plan_id: int, db: AsyncSession = Depends(get_db),
                             user: User = Depends(current_user)):
    """删除教案"""
    plan = await _get_plan(db, plan_id, user.id)
    if _plan_status(plan) == "processing":
        raise HTTPException(status_code=409, detail="教案正在生成中，完成后才能删除")
    await db.delete(plan)
    await db.commit()
    return {"message": "教案已删除", "id": plan_id}


@router.post("/lessonplan/{plan_id}/regenerate")
async def regenerate_lesson_plan(plan_id: int, wait: bool = Form(False), db: AsyncSession = Depends(get_db),
                                 user: User = Depends(current_user)):
    """按原参数重新生成教案（覆盖原内容）"""
    plan = await _get_plan(db, plan_id, user.id)
    if _plan_status(plan) == "processing":
        raise HTTPException(status_code=409, detail="教案正在生成中")
    plan.status = "processing"
    plan.progress_stage = "排队中"
    plan.error_message = None
    await db.commit()
    return await _run_or_spawn(db, plan, wait)


# ==================== 导出 ====================

def _markdown_to_docx(title: str, markdown: str) -> bytes:
    """把 Markdown 教案粗略转换为 Word 文档（标题 / 列表 / 表格 / 段落 / 粗体）"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def add_runs(paragraph, text: str) -> None:
        text = re.sub(r"\$\$?(.+?)\$\$?", r"\1", text)  # 公式保留原文
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                paragraph.add_run(part[2:-2]).bold = True
            elif part:
                paragraph.add_run(part.replace("`", ""))

    has_h1 = bool(re.search(r"^#\s", markdown, re.MULTILINE))
    if not has_h1:
        doc.add_heading(title, level=0)

    table_rows = []

    def flush_table():
        if not table_rows:
            return
        cols = max(len(r) for r in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell
        table_rows.clear()

    for raw in markdown.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            table_rows.append(cells)
            continue
        flush_table()
        if not stripped or stripped in ("---", "***"):
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            doc.add_heading(heading.group(2).replace("**", ""), level=0 if level == 1 else min(level - 1, 4))
            continue
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        if bullet:
            add_runs(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            continue
        numbered = re.match(r"^\d+[.)、]\s*(.*)$", stripped)
        if numbered:
            add_runs(doc.add_paragraph(style="List Number"), numbered.group(1))
            continue
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip("> "))
            continue
        add_runs(doc.add_paragraph(), stripped)
    flush_table()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _attachment_headers(filename: str) -> dict:
    ascii_name = re.sub(r"[^A-Za-z0-9._-]", "_", filename) or "lesson_plan"
    return {"Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"}


@router.get("/lessonplan/{plan_id}/export")
async def export_lesson_plan(
    plan_id: int,
    format: str = Query("md", pattern="^(md|txt|docx)$"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """导出教案：md / txt / docx"""
    plan = await _get_plan(db, plan_id, user.id)
    if not (plan.content or "").strip():
        raise HTTPException(status_code=400, detail="教案内容为空，无法导出")
    safe_title = re.sub(r'[\\/:*?"<>|\r\n]+', "_", plan.title).strip() or "教案"
    if format == "docx":
        data = _markdown_to_docx(plan.title, plan.content)
        return Response(
            content=data,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers=_attachment_headers(f"{safe_title}.docx"),
        )
    media = "text/markdown; charset=utf-8" if format == "md" else "text/plain; charset=utf-8"
    return Response(
        content=plan.content.encode("utf-8"),
        media_type=media,
        headers=_attachment_headers(f"{safe_title}.{format}"),
    )
