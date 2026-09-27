"""上传文件的安全保存与文本提取（图片 OCR / PDF / Word / TXT / MD）"""
import asyncio
import os
import tempfile
import uuid
from typing import Dict, List, Optional, Tuple

import aiofiles
from fastapi import HTTPException, UploadFile

from config import settings

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "bmp", "webp"}
DOC_EXTENSIONS = {"pdf", "docx", "txt", "md"}
ALLOWED_UPLOAD_EXTENSIONS = IMAGE_EXTENSIONS | DOC_EXTENSIONS

# 同时进行的 OCR 调用上限（真实模型单次 20-60 秒）
OCR_CONCURRENCY = 3

_CHUNK = 1024 * 1024


class FileParseError(Exception):
    """文件无法解析（非模型错误）"""


def ensure_upload_dir() -> str:
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    return settings.UPLOAD_DIR


def get_extension(filename: Optional[str]) -> str:
    name = os.path.basename(filename or "")
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def safe_display_name(filename: Optional[str]) -> str:
    """仅保留原始文件名的 basename，用于展示（不参与存储路径）"""
    name = os.path.basename((filename or "").replace("\\", "/"))
    return name[:200] or "未命名文件"


def max_size_label() -> str:
    return f"{settings.MAX_FILE_SIZE / 1024 / 1024:.0f}MB"


MSG_CONTENT_MISMATCH = "文件内容与格式不符"


def content_matches_extension(head: bytes, ext: str) -> bool:
    """
    文件头（magic bytes）校验：扩展名说是什么，内容就必须真是什么。
    图片 / PDF / Word 按文件头判断；txt / md 必须是文本（不含 NUL 字节，能按 UTF-8 或 GBK 解码）。
    """
    ext = (ext or "").lower()
    if ext in ("jpg", "jpeg"):
        return head.startswith(b"\xff\xd8\xff")
    if ext == "png":
        return head.startswith(b"\x89PNG\r\n\x1a\n")
    if ext == "gif":
        return head.startswith((b"GIF87a", b"GIF89a"))
    if ext == "bmp":
        return head.startswith(b"BM")
    if ext == "webp":
        return len(head) >= 12 and head[:4] == b"RIFF" and head[8:12] == b"WEBP"
    if ext == "pdf":
        return head[:1024].lstrip(b"\xef\xbb\xbf\r\n\t ").startswith(b"%PDF-")
    if ext == "docx":
        return head.startswith(b"PK\x03\x04")
    if ext in ("txt", "md", "csv"):
        if b"\x00" in head:
            return False
        sample = head[:4096]
        for encoding in ("utf-8", "gbk"):
            try:
                sample.decode(encoding)
                return True
            except UnicodeDecodeError as e:
                # 截断在多字节字符中间时，只要错误出现在末尾几个字节内就算通过
                if e.start >= len(sample) - 3:
                    return True
        return False
    return False


def validate_upload(file: UploadFile) -> str:
    """校验扩展名，返回小写扩展名；不合法时抛 400"""
    ext = get_extension(file.filename)
    if ext not in ALLOWED_UPLOAD_EXTENSIONS:
        allowed = "、".join(sorted(ALLOWED_UPLOAD_EXTENSIONS))
        raise HTTPException(
            status_code=400,
            detail=f"不支持的文件类型：{safe_display_name(file.filename)}（支持 {allowed}）",
        )
    return ext


async def save_upload(file: UploadFile, ext: Optional[str] = None) -> Dict:
    """
    以 uuid 文件名安全保存上传文件，并限制大小。
    返回 {"path", "original_name", "ext", "size"}；超限抛 413。
    """
    ext = ext or validate_upload(file)
    upload_dir = ensure_upload_dir()
    stored_name = f"{uuid.uuid4().hex}.{ext}"
    filepath = os.path.join(upload_dir, stored_name)

    size = 0
    first = True
    try:
        async with aiofiles.open(filepath, "wb") as out_file:
            while True:
                chunk = await file.read(_CHUNK)
                if not chunk:
                    break
                if first:
                    first = False
                    if not content_matches_extension(chunk, ext):
                        raise HTTPException(
                            status_code=400,
                            detail=MSG_CONTENT_MISMATCH,
                        )
                size += len(chunk)
                if size > settings.MAX_FILE_SIZE:
                    raise HTTPException(
                        status_code=413,
                        detail=f"文件过大：{safe_display_name(file.filename)}（单个文件不超过 {max_size_label()}）",
                    )
                await out_file.write(chunk)
    except HTTPException:
        _silent_remove(filepath)
        raise

    if size == 0:
        _silent_remove(filepath)
        raise HTTPException(status_code=400, detail=f"文件为空：{safe_display_name(file.filename)}")

    return {
        "path": filepath,
        "original_name": safe_display_name(file.filename),
        "ext": ext,
        "size": size,
    }


async def save_uploads(files: List[UploadFile]) -> List[Dict]:
    """先整体校验扩展名，再逐个保存；任何一个失败都会清理已保存的文件"""
    if not files:
        raise HTTPException(status_code=400, detail="请至少上传一个文件")
    limit = settings.MAX_FILES_PER_SUBMISSION
    if len(files) > limit:
        raise HTTPException(status_code=400, detail=f"一份作业最多上传 {limit} 个文件（当前 {len(files)} 个）")
    exts = [validate_upload(f) for f in files]
    saved: List[Dict] = []
    try:
        for f, ext in zip(files, exts):
            saved.append(await save_upload(f, ext))
    except Exception:
        for item in saved:
            _silent_remove(item["path"])
        raise
    return saved


def _silent_remove(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


# ==================== 文本提取 ====================

async def _read_text_file(filepath: str) -> str:
    async with aiofiles.open(filepath, "rb") as f:
        raw = await f.read()
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _read_docx(filepath: str) -> str:
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover
        raise FileParseError("服务器未安装 python-docx，无法解析 Word 文件") from e
    try:
        doc = Document(filepath)
    except Exception as e:
        raise FileParseError(f"Word 文件解析失败：{e}") from e
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_extract(filepath: str) -> Tuple[List[str], List[str]]:
    """
    返回 (每页文本, 需要 OCR 的页面图片路径)。
    优先 PyMuPDF（可把无文字的扫描页渲染成图片再 OCR），其次 pypdf（仅文本）。
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        fitz = None

    if fitz is not None:
        texts: List[str] = []
        images: List[str] = []
        try:
            doc = fitz.open(filepath)
        except Exception as e:
            raise FileParseError(f"PDF 文件解析失败：{e}") from e
        try:
            for page in doc:
                text = page.get_text().strip()
                if len(text) >= 10:
                    texts.append(text)
                else:
                    pix = page.get_pixmap(dpi=150)
                    tmp = os.path.join(tempfile.gettempdir(), f"pdfpage_{uuid.uuid4().hex}.png")
                    pix.save(tmp)
                    images.append(tmp)
                    texts.append("")
        finally:
            doc.close()
        return texts, images

    try:
        from pypdf import PdfReader
    except ImportError:
        PdfReader = None
    if PdfReader is not None:
        try:
            reader = PdfReader(filepath)
            texts = [(page.extract_text() or "").strip() for page in reader.pages]
        except Exception as e:
            raise FileParseError(f"PDF 文件解析失败：{e}") from e
        if any(len(t) >= 10 for t in texts):
            return texts, []
        raise FileParseError("该 PDF 为扫描版（无文字层），服务器缺少 PyMuPDF 无法识别，请转为图片后上传")

    raise FileParseError("服务器未安装 PDF 解析组件（pip install pymupdf），请将 PDF 转为图片后上传")


async def extract_text(
    filepath: str,
    ext: str,
    *,
    semaphore: Optional[asyncio.Semaphore] = None,
    task_meta: Optional[Dict] = None,
) -> str:
    """
    提取单个文件的文本。图片 / 扫描版 PDF 页走模型 OCR（受 semaphore 限流）。
    模型错误（llm.LLMError）原样抛出；文件本身无法解析时抛 FileParseError。
    """
    from ai_client import ai_client

    sem = semaphore or asyncio.Semaphore(OCR_CONCURRENCY)

    async def _ocr(path: str, meta: Dict) -> str:
        async with sem:
            return await ai_client.ocr_image(path, task_meta=meta)

    ext = (ext or "").lower()
    meta = dict(task_meta or {})

    if ext in IMAGE_EXTENSIONS:
        return await _ocr(filepath, meta)

    if ext in ("txt", "md"):
        return await _read_text_file(filepath)

    if ext == "docx":
        return await asyncio.to_thread(_read_docx, filepath)

    if ext == "pdf":
        texts, images = await asyncio.to_thread(_pdf_extract, filepath)
        if images:
            try:
                ocr_texts = await asyncio.gather(*[
                    _ocr(img, {**meta, "pdf_page_image": i}) for i, img in enumerate(images)
                ])
            finally:
                for img in images:
                    _silent_remove(img)
            it = iter(ocr_texts)
            texts = [t if t else next(it, "") for t in texts]
        return "\n\n".join(t for t in texts if t)

    raise FileParseError(f"不支持的文件类型：.{ext}")
